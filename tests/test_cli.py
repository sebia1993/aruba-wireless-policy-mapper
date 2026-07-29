from pathlib import Path

from openpyxl import Workbook, load_workbook
import json

import wlc_role_acl_collector.cli as cli
from wlc_role_acl_collector.cli import main
from wlc_role_acl_collector.diagnostic_mode import DiagnosticRun
from wlc_role_acl_collector.interactive import prompt_controller_targets
from wlc_role_acl_collector.models import (
    CollectionResult,
    CommandOutput,
    Controller,
    ControllerCredentials,
    ControllerTarget,
    ParsedController,
)


def test_cli_collect_offline(tmp_path):
    controllers = tmp_path / "controllers.csv"
    controllers.write_text(
        "name,host,protocol,port,device_type,username_env,password_env,enable_password_env\n"
        "sample_controller,192.0.2.10,ssh,22,aruba_os,,,\n",
        encoding="utf-8",
    )
    fixture_root = Path(__file__).parent / "fixtures"
    role_networks = tmp_path / "role_networks.xlsx"
    _write_role_networks(role_networks)

    exit_code = main(
        [
            "collect",
            "--controllers",
            str(controllers),
            "--offline-raw-dir",
            str(fixture_root),
            "--output-dir",
            str(tmp_path / "outputs"),
            "--role-networks",
            str(role_networks),
        ]
    )

    assert exit_code == 0
    reports = list((tmp_path / "outputs").glob("*/ssid_role_acl_report.xlsx"))
    assert len(reports) == 1
    workbook = load_workbook(reports[0], read_only=True)
    assert "Local_Role_Networks" not in workbook.sheetnames


def test_cli_can_explicitly_export_local_role_networks(tmp_path):
    controllers = tmp_path / "controllers.csv"
    controllers.write_text(
        "name,host,protocol,port,device_type,username_env,password_env,enable_password_env\n"
        "sample_controller,192.0.2.10,ssh,22,aruba_os,,,\n",
        encoding="utf-8",
    )
    fixture_root = Path(__file__).parent / "fixtures"
    role_networks = tmp_path / "role_networks.xlsx"
    _write_role_networks(role_networks)

    exit_code = main(
        [
            "collect",
            "--controllers",
            str(controllers),
            "--offline-raw-dir",
            str(fixture_root),
            "--output-dir",
            str(tmp_path / "outputs"),
            "--role-networks",
            str(role_networks),
            "--export-local-role-networks",
        ]
    )

    assert exit_code == 0
    reports = list((tmp_path / "outputs").glob("*/ssid_role_acl_report.xlsx"))
    assert len(reports) == 1
    workbook = load_workbook(reports[0], read_only=True)
    assert "Local_Role_Networks" in workbook.sheetnames


def test_cli_interactive_defaults_to_ssh(monkeypatch, tmp_path):
    prompts = iter(
        [
            "10.10.10.10",
            "",
            "",
            "",
            "admin",
            "",
        ]
    )
    captured = {}

    monkeypatch.setattr(
        cli,
        "prompt_controller_targets",
        lambda: prompt_controller_targets(
            input_func=lambda _prompt: next(prompts),
            password_func=lambda _prompt: "secret",
        ),
    )

    def fake_collect(controller, *, timeout, credentials):
        captured["controller"] = controller
        captured["credentials"] = credentials
        return CollectionResult(
            controller=controller,
            commands=[
                CommandOutput(
                    command_id="configuration_effective",
                    command="show configuration effective",
                    output="show configuration effective\n",
                )
            ],
        )

    monkeypatch.setattr(cli, "collect_from_controller", fake_collect)

    exit_code = main(["collect", "--output-dir", str(tmp_path / "outputs")])

    assert exit_code == 0
    assert captured["controller"].host == "10.10.10.10"
    assert captured["controller"].name == "wlc-10.10.10.10"
    assert captured["controller"].protocol == "ssh"
    assert captured["controller"].port == 22
    assert captured["controller"].device_type == "aruba_os"
    assert captured["credentials"].username == "admin"


def test_cli_returns_failure_when_required_configuration_was_not_collected(monkeypatch, tmp_path, capsys):
    controller = Controller(name="failed-wlc", host="192.0.2.10")
    target = ControllerTarget(
        controller=controller,
        credentials=ControllerCredentials(username="admin", password="secret"),
    )
    monkeypatch.setattr(cli, "_resolve_targets", lambda _args: [target])
    monkeypatch.setattr(
        cli,
        "collect_from_controller",
        lambda *_args, **_kwargs: CollectionResult(
            controller=controller,
            commands=[
                CommandOutput(
                    command_id="connect",
                    command="connect",
                    success=False,
                    error="authentication failed",
                )
            ],
        ),
    )

    exit_code = main(["collect", "--output-dir", str(tmp_path / "outputs")])

    assert exit_code == cli.COLLECT_EXIT_FAILED
    assert "failed-wlc" in capsys.readouterr().err


def test_cli_returns_partial_when_optional_command_failed(monkeypatch, tmp_path, capsys):
    controller = Controller(name="partial-wlc", host="192.0.2.10")
    target = ControllerTarget(
        controller=controller,
        credentials=ControllerCredentials(username="admin", password="secret"),
    )
    monkeypatch.setattr(cli, "_resolve_targets", lambda _args: [target])
    monkeypatch.setattr(
        cli,
        "collect_from_controller",
        lambda *_args, **_kwargs: CollectionResult(
            controller=controller,
            commands=[
                CommandOutput(
                    command_id="configuration_effective",
                    command="show configuration effective",
                    output="show configuration effective\n",
                ),
                CommandOutput(
                    command_id="user_table",
                    command="show user-table",
                    success=False,
                    error="command timed out",
                ),
            ],
        ),
    )

    exit_code = main(["collect", "--output-dir", str(tmp_path / "outputs")])

    assert exit_code == cli.COLLECT_EXIT_PARTIAL
    assert "partial-wlc:user_table" in capsys.readouterr().err


def test_cli_rejects_timeout_outside_supported_range(capsys):
    exit_code = main(["collect", "--timeout", "1"])

    assert exit_code == cli.COLLECT_EXIT_INPUT_ERROR
    assert "5에서 600" in capsys.readouterr().err


def test_cli_invalid_controller_file_returns_input_error_without_traceback(tmp_path, capsys):
    exit_code = main(
        [
            "collect",
            "--controllers",
            str(tmp_path / "missing.csv"),
            "--output-dir",
            str(tmp_path / "outputs"),
        ]
    )

    error = capsys.readouterr().err
    assert exit_code == cli.COLLECT_EXIT_INPUT_ERROR
    assert "WLC 대상 설정 오류" in error
    assert "Traceback" not in error


def test_cli_collection_exception_does_not_stop_remaining_wlc(monkeypatch, tmp_path, capsys):
    bad = Controller(name="bad-wlc", host="192.0.2.10")
    good = Controller(name="good-wlc", host="192.0.2.11")
    targets = [
        ControllerTarget(bad, ControllerCredentials(username="admin", password="secret")),
        ControllerTarget(good, ControllerCredentials(username="admin", password="secret")),
    ]
    calls = []
    monkeypatch.setattr(cli, "_resolve_targets", lambda _args: targets)

    def fake_collect(controller, **_kwargs):
        calls.append(controller.name)
        if controller.name == "bad-wlc":
            raise RuntimeError("Authentication failed: bad password")
        return CollectionResult(
            controller=controller,
            commands=[
                CommandOutput(
                    command_id="configuration_effective",
                    command="show configuration effective",
                    output="show configuration effective\n",
                )
            ],
        )

    monkeypatch.setattr(cli, "collect_from_controller", fake_collect)

    exit_code = main(["collect", "--output-dir", str(tmp_path / "outputs")])

    error = capsys.readouterr().err
    assert exit_code == cli.COLLECT_EXIT_FAILED
    assert calls == ["bad-wlc", "good-wlc"]
    assert "WLC-AUTH-001" in error
    assert "Traceback" not in error
    assert len(list((tmp_path / "outputs").glob("*/ssid_role_acl_report.html"))) == 1
    assert len(list((tmp_path / "outputs").glob("*/raw/*.txt"))) == 2


def test_cli_parser_exception_isolated_to_one_controller(monkeypatch, tmp_path, capsys):
    first = Controller(name="first-wlc", host="192.0.2.10")
    second = Controller(name="second-wlc", host="192.0.2.11")
    targets = [
        ControllerTarget(first, ControllerCredentials(username="admin", password="secret")),
        ControllerTarget(second, ControllerCredentials(username="admin", password="secret")),
    ]
    monkeypatch.setattr(cli, "_resolve_targets", lambda _args: targets)
    monkeypatch.setattr(
        cli,
        "collect_from_controller",
        lambda controller, **_kwargs: CollectionResult(
            controller=controller,
            commands=[
                CommandOutput(
                    command_id="configuration_effective",
                    command="show configuration effective",
                    output="show configuration effective\n",
                )
            ],
        ),
    )

    def fake_parse(results):
        result = results[0]
        if result.controller.name == "first-wlc":
            raise ValueError("unsupported private output")
        return [ParsedController(controller=result.controller)]

    monkeypatch.setattr(cli, "build_parsed_controllers", fake_parse)

    exit_code = main(["collect", "--output-dir", str(tmp_path / "outputs")])

    error = capsys.readouterr().err
    run_dir = next((tmp_path / "outputs").iterdir())
    status = json.loads((run_dir / "report_status.json").read_text(encoding="utf-8"))
    assert exit_code == cli.COLLECT_EXIT_FAILED
    assert "WLC-PRS-001" in error
    assert "Traceback" not in error
    assert status["status"] == "completed"
    assert status["collection_status"] == "partial"
    assert (run_dir / "ssid_role_acl_report.html").exists()


def test_cli_report_error_returns_controlled_failure(monkeypatch, tmp_path, capsys):
    controller = Controller(name="wlc", host="192.0.2.10")
    target = ControllerTarget(
        controller=controller,
        credentials=ControllerCredentials(username="admin", password="secret"),
    )
    monkeypatch.setattr(cli, "_resolve_targets", lambda _args: [target])
    monkeypatch.setattr(
        cli,
        "collect_from_controller",
        lambda *_args, **_kwargs: CollectionResult(
            controller=controller,
            commands=[
                CommandOutput(
                    command_id="configuration_effective",
                    command="show configuration effective",
                    output="show configuration effective\n",
                )
            ],
        ),
    )
    monkeypatch.setattr(
        cli,
        "write_reports",
        lambda **_kwargs: (_ for _ in ()).throw(OSError("disk full")),
    )

    exit_code = main(["collect", "--output-dir", str(tmp_path / "outputs")])

    error = capsys.readouterr().err
    assert exit_code == cli.COLLECT_EXIT_FAILED
    assert "WLC-RPT-002" in error
    assert "Traceback" not in error


def test_cli_output_path_failure_returns_controlled_error(tmp_path, capsys):
    blocked_output = tmp_path / "output-is-a-file"
    blocked_output.write_text("not a directory", encoding="utf-8")
    controllers = tmp_path / "controllers.csv"
    controllers.write_text(
        "name,host,protocol,port,device_type,username_env,password_env,enable_password_env\n"
        "sample_controller,192.0.2.10,ssh,22,aruba_os,,,\n",
        encoding="utf-8",
    )

    exit_code = main(
        [
            "collect",
            "--controllers",
            str(controllers),
            "--offline-raw-dir",
            str(Path(__file__).parent / "fixtures"),
            "--output-dir",
            str(blocked_output),
        ]
    )

    error = capsys.readouterr().err
    assert exit_code == cli.COLLECT_EXIT_FAILED
    assert "WLC-RPT-001" in error
    assert "Traceback" not in error


def test_cli_diagnostic_exception_does_not_stop_remaining_wlc(monkeypatch, tmp_path, capsys):
    first = Controller(name="first-wlc", host="192.0.2.10")
    second = Controller(name="second-wlc", host="192.0.2.11")
    targets = [
        ControllerTarget(first, ControllerCredentials(username="admin", password="secret")),
        ControllerTarget(second, ControllerCredentials(username="admin", password="secret")),
    ]
    calls = []
    monkeypatch.setattr(cli, "_resolve_targets", lambda _args: targets)

    def fake_diagnostic(target, **_kwargs):
        calls.append(target.controller.name)
        if target.controller.name == "first-wlc":
            raise RuntimeError("Authentication failed: bad password")
        return DiagnosticRun(
            run_dir=tmp_path / "diagnostic",
            primary_code="OK",
            report_paths={
                "json": tmp_path / "diagnostic_summary.json",
                "html": tmp_path / "diagnostic_summary.html",
            },
            events=[],
        )

    monkeypatch.setattr(cli, "run_diagnostic", fake_diagnostic)

    exit_code = main(["diagnose", "--output-dir", str(tmp_path / "outputs")])

    error = capsys.readouterr().err
    assert exit_code == cli.COLLECT_EXIT_FAILED
    assert calls == ["first-wlc", "second-wlc"]
    assert "WLC-AUTH-001" in error
    assert "Traceback" not in error


def _write_role_networks(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["role", "network", "subnet_mask"])
    worksheet.append(["guest-logon", "10.30.0.0", "255.255.255.0"])
    workbook.save(path)
